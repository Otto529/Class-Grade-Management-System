# -*- coding: utf-8 -*-
"""
excel 的导入导出都在这。
名单：学号 姓名 性别 班级
成绩：学号 姓名 班级 语文 数学 ...，第一行是表头
"""
import io
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core import database as db
from core import repository as repo

# 表头样式
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def _autofit(ws):
    # 列宽按内容大概调一下
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=4)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width * 2 + 4, 24)


# 名单模板
def student_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "学生名单"
    ws.append(["学号", "姓名", "性别", "班级"])
    ws.append(["20240101", "张三", "男", "高三(1)班"])
    ws.append(["20240102", "李四", "女", "高三(1)班"])
    _style_header(ws, 4)
    _autofit(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# 成绩模板，学生名单直接带出来，老师只填分
def score_template():
    subjects = repo.list_subjects()
    wb = Workbook()
    ws = wb.active
    ws.title = "成绩录入"
    headers = ["学号", "姓名", "班级"] + [s["name"] for s in subjects]
    ws.append(headers)
    students = db.query(
        """SELECT st.student_no, st.name, c.class_name
           FROM students st JOIN classes c ON c.id=st.class_id
           WHERE st.status=1 ORDER BY st.student_no"""
    )
    for st in students:
        ws.append([st["student_no"], st["name"], st["class_name"]] + [None] * len(subjects))
    _style_header(ws, len(headers))
    _autofit(ws)

    ws2 = wb.create_sheet("填写说明")
    ws2.append(["说明"])
    ws2.append(["1. 表头别动，学号要和系统里一致"])
    ws2.append(["2. 缺考空着就行"])
    ws2.append(["3. 分数别超过满分（语数英150，其余100）"])
    ws2.column_dimensions["A"].width = 60

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _read_sheet(file_stream):
    wb = load_workbook(file_stream, data_only=True)
    return wb.active


# 导入名单，已有的学号按更新处理
def import_students(file_stream):
    ws = _read_sheet(file_stream)
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    classes = {c["class_name"]: c["id"] for c in repo.list_classes()}
    existing = {r["student_no"] for r in db.query("SELECT student_no FROM students")}

    records, errors = [], []
    for i, row in enumerate(rows, start=2):
        if not row or row[0] is None:
            continue
        student_no, name, gender, class_name = (list(row) + [None] * 4)[:4]
        student_no = str(student_no).strip()
        name = str(name).strip() if name else None
        gender = str(gender).strip() if gender else "男"
        class_name = str(class_name).strip() if class_name else None
        if not name or class_name not in classes:
            errors.append(f"第{i}行：姓名为空或班级「{class_name}」不存在")
            continue
        if gender not in ("男", "女"):
            gender = "男"
        if student_no in existing:
            db.execute(
                "UPDATE students SET name=%s,gender=%s,class_id=%s,status=1 WHERE student_no=%s",
                (name, gender, classes[class_name], student_no),
            )
            continue
        records.append((student_no, name, gender, classes[class_name]))
        existing.add(student_no)

    if records:
        db.executemany(
            "INSERT INTO students(student_no,name,gender,class_id) VALUES(%s,%s,%s,%s)",
            records,
        )
    return len(records), errors


# 导入成绩
def import_scores(file_stream, exam_id):
    ws = _read_sheet(file_stream)
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return 0, ["文件为空"]
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]

    subjects = {s["name"]: (s["id"], float(s["full_score"])) for s in repo.list_subjects()}
    subject_cols = []
    for idx, h in enumerate(headers):
        if h in subjects:
            subject_cols.append((idx, h))
    if not subject_cols:
        return 0, ["表头里没认出来科目列"]

    stu_rows = db.query("SELECT id, student_no FROM students WHERE status=1")
    stu_map = {s["student_no"]: s["id"] for s in stu_rows}

    records, errors, seen = [], [], set()
    for i, row in enumerate(rows[1:], start=2):
        if not row or row[0] is None:
            continue
        student_no = str(row[0]).strip()
        student_id = stu_map.get(student_no)
        if not student_id:
            errors.append(f"第{i}行：学号「{student_no}」不在系统里")
            continue
        for col_idx, subject_name in subject_cols:
            val = row[col_idx] if col_idx < len(row) else None
            if val is None or str(val).strip() == "":
                continue
            try:
                score = float(val)
            except (TypeError, ValueError):
                errors.append(f"第{i}行 {subject_name}：「{val}」不是数字")
                continue
            subj_id, full = subjects[subject_name]
            if score < 0 or score > full:
                errors.append(f"第{i}行 {subject_name}：{score} 超出 0~{full}")
                continue
            key = (exam_id, student_id, subj_id)
            if key in seen:
                continue
            seen.add(key)
            records.append((exam_id, student_id, subj_id, score))

    if records:
        repo.bulk_upsert_scores(exam_id, records)
    return len(records), errors


# 导出成绩册，两个 sheet：成绩册 + 科目统计
def export_report(exam_id, class_id=None):
    subjects, students = repo.score_report(exam_id, class_id)
    exam = db.query_one("SELECT * FROM exams WHERE id=%s", (exam_id,))
    stats = repo.subject_stats(exam_id, class_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "成绩册"
    headers = ["学号", "姓名", "班级"] + [s["name"] for s in subjects] + ["总分", "平均分", "排名"]
    ws.append(headers)

    # 总分和排名在这边算好再写进去
    totals = []
    for st in students:
        vals = [st["scores"][s["name"]] for s in subjects]
        valid = [v for v in vals if v is not None]
        total = round(sum(valid), 1) if valid else None
        avg = round(total / len(valid), 2) if valid else None
        totals.append((st, vals, total, avg))
    rank_order = sorted([t for t in totals if t[2] is not None],
                        key=lambda x: x[2], reverse=True)
    rank_map, last_total, last_rank = {}, None, 0
    for i, t in enumerate(rank_order, start=1):
        if t[2] != last_total:
            last_rank, last_total = i, t[2]
        rank_map[t[0]["student_no"]] = last_rank

    for st, vals, total, avg in totals:
        ws.append([st["student_no"], st["name"], st["class_name"]] + vals +
                  [total, avg, rank_map.get(st["student_no"])])
    _style_header(ws, len(headers))
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in r:
            cell.alignment = CENTER
            cell.border = BORDER
    _autofit(ws)
    ws.freeze_panes = "D2"

    ws2 = wb.create_sheet("科目统计")
    ws2.append(["科目", "参考人数", "平均分", "最高分", "最低分", "及格率%", "优秀率%"])
    for s in stats:
        ws2.append([s["name"], s["cnt"], s["avg_score"], s["max_score"],
                    s["min_score"], s["pass_rate"], s["excellent_rate"]])
    _style_header(ws2, 7)
    _autofit(ws2)

    fname = f"{exam['name']}_成绩册_{datetime.now():%Y%m%d}.xlsx"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return fname, buf
